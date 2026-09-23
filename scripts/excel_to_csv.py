"""Export every worksheet from the IEK and Systeme electric Excel folders to CSV.

Usage: python scripts/excel_to_csv.py [DATASETS_DIR]
Requires: openpyxl (listed in requirements.txt).
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCES = ("IEK", "Systeme electric")
INVALID_FILENAME_CHARS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def safe_name(name: str) -> str:
    cleaned = INVALID_FILENAME_CHARS.sub("_", name).rstrip(" .")
    return cleaned or "Sheet"


def csv_value(value: object) -> object:
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat(sep=" ") if isinstance(value, dt.datetime) else value.isoformat()
    if isinstance(value, dt.timedelta):
        return str(value)
    return value


def convert_folder(source: Path) -> tuple[int, int]:
    if not source.is_dir():
        raise FileNotFoundError(f"Source folder does not exist: {source}")

    files = sorted(
        path for path in source.rglob("*")
        if path.is_file()
        and path.suffix.lower() in {".xlsx", ".xlsm"}
        and not path.name.startswith("~$")
    )
    if not files:
        raise FileNotFoundError(f"No .xlsx or .xlsm files found in {source}")

    destination = source.with_name(source.name + "_CSV")
    used: set[str] = set()
    exported = 0
    for book_path in files:
        workbook = load_workbook(book_path, read_only=True, data_only=True)
        try:
            relative_dir = book_path.parent.relative_to(source)
            for sheet in workbook.worksheets:
                base = book_path.stem
                if len(workbook.worksheets) > 1:
                    base += "__" + sheet.title
                candidate = destination / relative_dir / (safe_name(base) + ".csv")
                key = str(candidate).casefold()
                if key in used:
                    raise ValueError(f"CSV filename collision: {candidate}")
                used.add(key)
                candidate.parent.mkdir(parents=True, exist_ok=True)
                with candidate.open("w", encoding="utf-8-sig", newline="") as output:
                    writer = csv.writer(output)
                    for row in sheet.iter_rows(values_only=True):
                        writer.writerow(csv_value(value) for value in row)
                exported += 1
                print(f"{book_path.relative_to(source)} [{sheet.title}] -> {candidate.relative_to(destination)}")
        finally:
            workbook.close()
    return len(files), exported


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "datasets_dir", type=Path, nargs="?",
        default=Path(__file__).resolve().parents[1] / "datasets",
        help="Directory containing IEK and Systeme electric (default: repository datasets/)",
    )
    args = parser.parse_args()
    for name in SOURCES:
        files, sheets = convert_folder(args.datasets_dir / name)
        print(f"{name}: {files} workbooks, {sheets} CSV files")


if __name__ == "__main__":
    main()
